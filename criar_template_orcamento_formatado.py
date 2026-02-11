#!/usr/bin/env python3
"""
Gera template de orçamento DOCX formatado com as chaves mapeadas
Baseado no estilo do contrato original
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_page_border(doc):
    """Adiciona borda na página"""
    sectPr = doc.sections[0]._sectPr
    pgBorders = OxmlElement('w:pgBorders')
    pgBorders.set(qn('w:offsetFrom'), 'page')
    
    for border_name in ('top', 'left', 'bottom', 'right'):
        border_el = OxmlElement(f'w:{border_name}')
        border_el.set(qn('w:val'), 'single')
        border_el.set(qn('w:sz'), '12')
        border_el.set(qn('w:space'), '24')
        border_el.set(qn('w:color'), 'F97316')
        pgBorders.append(border_el)
    
    sectPr.append(pgBorders)

def criar_template_orcamento():
    """Cria template de orçamento formatado"""
    doc = Document()
    
    # Configurar margens
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
    
    # Adicionar borda
    add_page_border(doc)
    
    # CABEÇALHO
    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run('PROPOSTA COMERCIAL')
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(249, 115, 22)  # Laranja
    
    # Número e Data
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Nº {{NUMERO_ORCAMENTO}} | {{DATA_ORCAMENTO}}')
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(107, 114, 128)
    
    doc.add_paragraph()  # Espaço
    
    # DADOS DO CLIENTE
    p = doc.add_paragraph()
    run = p.add_run('DADOS DO CLIENTE')
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(249, 115, 22)
    
    # Tabela Cliente
    table = doc.add_table(rows=5, cols=2)
    table.style = 'Light Grid Accent 1'
    
    dados_cliente = [
        ('Nome:', '{{NOME_CLIENTE}}'),
        ('CPF/CNPJ:', '{{CPF_CNPJ}}'),
        ('Telefone:', '{{TELEFONE}}'),
        ('E-mail:', '{{EMAIL}}'),
        ('Endereço:', '{{ENDERECO}}, {{CIDADE}} - {{ESTADO}}'),
    ]
    
    for i, (label, valor) in enumerate(dados_cliente):
        row = table.rows[i]
        row.cells[0].text = label
        row.cells[0].paragraphs[0].runs[0].font.bold = True
        row.cells[1].text = valor
    
    doc.add_paragraph()  # Espaço
    
    # SISTEMA PROPOSTO
    p = doc.add_paragraph()
    run = p.add_run('SISTEMA FOTOVOLTAICO PROPOSTO')
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(249, 115, 22)
    
    # Tabela Sistema
    table = doc.add_table(rows=5, cols=2)
    table.style = 'Light Grid Accent 1'
    
    dados_sistema = [
        ('Potência Total:', '{{POTENCIA_KWP}} kWp'),
        ('Painéis Solares:', '{{QUANTIDADE_PAINEIS}}x {{MARCA_PAINEL}} {{POTENCIA_PAINEL}}W'),
        ('Inversor:', '{{QUANTIDADE_INVERSORES}}x {{MARCA_INVERSOR}} {{POTENCIA_INVERSOR_KW}}kW'),
        ('Estrutura:', '{{TIPO_ESTRUTURA}}'),
        ('Geração Mensal:', '{{GERACAO_MENSAL}} kWh/mês'),
    ]
    
    for i, (label, valor) in enumerate(dados_sistema):
        row = table.rows[i]
        row.cells[0].text = label
        row.cells[0].paragraphs[0].runs[0].font.bold = True
        row.cells[1].text = valor
    
    doc.add_paragraph()  # Espaço
    
    # COMPOSIÇÃO DO INVESTIMENTO
    p = doc.add_paragraph()
    run = p.add_run('COMPOSIÇÃO DO INVESTIMENTO')
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(249, 115, 22)
    
    # Tabela Valores
    table = doc.add_table(rows=7, cols=2)
    table.style = 'Light Grid Accent 1'
    
    dados_valores = [
        ('Kit Fotovoltaico:', '{{VALOR_KIT}}'),
        ('Estrutura de Fixação:', '{{VALOR_ESTRUTURA}}'),
        ('Materiais Elétricos:', '{{VALOR_MATERIAL_ELETRICO}}'),
        ('Projeto Elétrico:', '{{VALOR_PROJETO}}'),
        ('Instalação/Montagem:', '{{VALOR_MONTAGEM}}'),
        ('', ''),
        ('VALOR TOTAL:', '{{VALOR_FINAL}}'),
    ]
    
    for i, (label, valor) in enumerate(dados_valores):
        row = table.rows[i]
        row.cells[0].text = label
        row.cells[1].text = valor
        
        if i == 6:  # Última linha (total)
            row.cells[0].paragraphs[0].runs[0].font.bold = True
            row.cells[0].paragraphs[0].runs[0].font.size = Pt(12)
            row.cells[1].paragraphs[0].runs[0].font.bold = True
            row.cells[1].paragraphs[0].runs[0].font.size = Pt(12)
            row.cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(34, 197, 94)
    
    doc.add_paragraph()  # Espaço
    
    # CONDIÇÕES DE PAGAMENTO
    p = doc.add_paragraph()
    run = p.add_run('CONDIÇÕES DE PAGAMENTO')
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(249, 115, 22)
    
    p = doc.add_paragraph()
    run = p.add_run('Forma de Pagamento: ')
    run.font.bold = True
    run = p.add_run('{{FORMA_PAGAMENTO}}')
    run.font.size = Pt(12)
    
    doc.add_paragraph()  # Espaço
    
    # PREMISSAS TÉCNICAS
    p = doc.add_paragraph()
    run = p.add_run('PREMISSAS TÉCNICAS')
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(249, 115, 22)
    
    premissas = [
        f'• Horas de Sol Pico (HSP): {{{{HSP}}}} h/dia',
        f'• Perda do Sistema: {{{{PERDA_SISTEMA}}}}',
        f'• Geração Anual Estimada: {{{{GERACAO_ANUAL}}}} kWh',
        f'• Validade da Proposta: {{{{DATA_VALIDADE}}}}',
    ]
    
    for premissa in premissas:
        p = doc.add_paragraph(premissa)
        p.style = 'List Bullet'
    
    doc.add_paragraph()  # Espaço
    
    # OBSERVAÇÕES
    p = doc.add_paragraph()
    run = p.add_run('OBSERVAÇÕES IMPORTANTES')
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(249, 115, 22)
    
    obs = [
        'Todos os equipamentos são novos e possuem garantia de fábrica',
        'Instalação conforme normas ABNT NBR 16690, 5410 e 5419',
        'Projeto elétrico e homologação junto à concessionária inclusos',
        'Garantia de 12 meses para instalação e 25 anos para painéis',
    ]
    
    for ob in obs:
        p = doc.add_paragraph(f'• {ob}')
        p.style = 'List Bullet'
    
    # RODAPÉ
    doc.add_paragraph()
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('_' * 50)
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('{{NOME_VENDEDOR}}')
    run.font.bold = True
    run.font.size = Pt(11)
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('{{TELEFONE_VENDEDOR}} | {{EMAIL_VENDEDOR}}')
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(107, 114, 128)
    
    # Salvar
    doc.save('docs/templates/TEMPLATE_ORCAMENTO.docx')
    print('Template criado: docs/templates/TEMPLATE_ORCAMENTO.docx')

if __name__ == '__main__':
    criar_template_orcamento()
