"""
Script para converter PDF de proposta comercial em template DOCX com chaves
"""
import subprocess
import sys

# Instalar dependências
print("Instalando pdf2docx...")
subprocess.check_call([sys.executable, "-m", "pip", "install", "pdf2docx", "-q"])

from pdf2docx import Converter
from docx import Document
from docx.shared import Pt, RGBColor
import os

# Converter PDF para DOCX
print("Convertendo PDF para DOCX...")
pdf_path = r"d:\SunOps---SaaS\PROPOSTA COMERCIAL - MAB (1).pdf"
temp_docx = r"d:\SunOps---SaaS\temp_proposta.docx"

cv = Converter(pdf_path)
cv.convert(temp_docx)
cv.close()

print("PDF convertido! Adicionando chaves de template...")

# Abrir o documento convertido
doc = Document(temp_docx)

# Mapeamento de textos para chaves
replacements = {
    # Dados do Cliente
    "MAB": "{{cliente_nome}}",
    "Rua Exemplo, 123": "{{cliente_endereco}}",
    "Bairro Exemplo": "{{cliente_bairro}}",
    "Cidade - UF": "{{cliente_cidade}} - {{cliente_estado}}",
    "CEP 00000-000": "{{cliente_cep}}",
    "(00) 00000-0000": "{{cliente_telefone}}",
    "cliente@email.com": "{{cliente_email}}",
    
    # Dados do Sistema
    "10.000 kWh": "{{consumo_mensal}} kWh",
    "15,00 kWp": "{{potencia_sistema}} kWp",
    "30": "{{quantidade_paineis}}",
    "500 Wp": "{{potencia_painel}} Wp",
    "15 kW": "{{potencia_inversor}} kW",
    "1.350 kWh": "{{geracao_mensal}} kWh",
    "16.200 kWh": "{{geracao_anual}} kWh",
    
    # Dados Financeiros
    "R$ 75.000,00": "{{valor_total}}",
    "R$ 6.250,00": "{{valor_parcela}}",
    "12x": "{{quantidade_parcelas}}x",
    "25 anos": "{{vida_util_sistema}} anos",
    "R$ 1.500,00": "{{economia_mensal}}",
    "R$ 18.000,00": "{{economia_anual}}",
    "4,2 anos": "{{payback}} anos",
    
    # Dados da Empresa
    "SunOps": "{{empresa_nome}}",
    "CNPJ 00.000.000/0001-00": "{{empresa_cnpj}}",
    "Rua da Empresa, 456": "{{empresa_endereco}}",
    "(00) 0000-0000": "{{empresa_telefone}}",
    "contato@sunops.com": "{{empresa_email}}",
    "www.sunops.com": "{{empresa_site}}",
    
    # Dados Técnicos
    "5,5 horas": "{{hsp}} horas",
    "20%": "{{perdas_sistema}}%",
    "0,8%": "{{degradacao_anual}}%",
    
    # Datas
    "01/01/2024": "{{data_proposta}}",
    "30 dias": "{{validade_proposta}} dias",
}

# Substituir em parágrafos
for paragraph in doc.paragraphs:
    for old_text, new_text in replacements.items():
        if old_text in paragraph.text:
            # Preservar formatação
            for run in paragraph.runs:
                if old_text in run.text:
                    run.text = run.text.replace(old_text, new_text)

# Substituir em tabelas
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for old_text, new_text in replacements.items():
                    if old_text in paragraph.text:
                        for run in paragraph.runs:
                            if old_text in run.text:
                                run.text = run.text.replace(old_text, new_text)

# Salvar template final
output_path = r"d:\SunOps---SaaS\docs\templates\PROPOSTA_COMERCIAL_TEMPLATE.docx"
os.makedirs(os.path.dirname(output_path), exist_ok=True)
doc.save(output_path)

# Remover arquivo temporário
os.remove(temp_docx)

print(f"\n✅ Template criado com sucesso!")
print(f"📄 Arquivo: {output_path}")
print(f"\n🔑 Chaves adicionadas:")
for key in sorted(set(replacements.values())):
    print(f"   - {key}")
