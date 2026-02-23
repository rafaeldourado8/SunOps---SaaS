"""
Utilitário para geração de Propostas Comerciais a partir do template DOCX
"""
from docx import Document
from datetime import datetime
import os


class GeradorPropostaComercial:
    """Gera propostas comerciais em DOCX a partir de dados do orçamento"""
    
    TEMPLATE_PATH = 'docs/templates/PROPOSTA_COMERCIAL_TEMPLATE.docx'
    
    # Chaves que foram substituidas por valores fixos no template
    MARCAS_FIXAS = {
        'painel': 'CANADIAN SOLAR',
        'inversor': 'GROWATT'
    }
    
    def __init__(self, orcamento):
        self.orcamento = orcamento
        self.cliente = orcamento.cliente
        
    def _preparar_dados(self):
        """Prepara dicionário com todos os dados para substituição"""
        o = self.orcamento
        c = self.cliente
        
        # Calcular potencia do sistema
        potencia_sistema = (o.potencia_painel * o.quantidade_paineis) / 1000
        
        return {
            # Dados do Cliente
            '{{cliente_nome}}': c.nome or '',
            '{{cliente_endereco}}': c.endereco or '',
            '{{cliente_bairro}}': c.bairro or '',
            '{{cliente_cidade}}': c.cidade or '',
            '{{cliente_estado}}': c.estado or '',
            '{{cliente_cep}}': c.cep or '',
            '{{cliente_telefone}}': c.telefone or '',
            '{{cliente_email}}': c.email or '',
            
            # Dados do Sistema
            '{{consumo_mensal}}': '0',
            '{{potencia_sistema}}': f"{potencia_sistema:.2f}",
            '{{POTENCIA_TOTAL_KWP}}': f"{potencia_sistema:.2f}",
            '{{quantidade_paineis}}': str(o.quantidade_paineis),
            '{{potencia_painel}}': str(o.potencia_painel),
            '{{PAINEIS_POTENCIA}}': str(o.potencia_painel),
            '{{marca_painel}}': o.marca_painel,
            '{{potencia_inversor}}': str(o.potencia_inversor),
            '{{INVERSOR_POTENCIA}}': str(o.potencia_inversor),
            '{{marca_inversor}}': o.marca_inversor,
            '{{quantidade_inversores}}': str(o.quantidade_inversores),
            '{{geracao_mensal}}': '0',
            '{{geracao_anual}}': '0',
            
            # Dados Financeiros
            '{{valor_total}}': f"R$ {o.valor_final:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.'),
            '{{valor_parcela}}': f"R$ {o.valor_parcela:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.') if o.valor_parcela else 'R$ 0,00',
            '{{quantidade_parcelas}}': '1',
            '{{vida_util_sistema}}': '25',
            '{{economia_mensal}}': 'R$ 0,00',
            '{{economia_anual}}': 'R$ 0,00',
            '{{payback}}': '0',
            
            # Dados Técnicos
            '{{hsp}}': '5,5',
            '{{perdas_sistema}}': '20',
            '{{degradacao_anual}}': '0,8',
            
            # Dados da Empresa
            '{{empresa_nome}}': os.getenv('EMPRESA_NOME', 'SunOps Energia Solar'),
            '{{empresa_cnpj}}': os.getenv('EMPRESA_CNPJ', '00.000.000/0001-00'),
            '{{empresa_endereco}}': os.getenv('EMPRESA_ENDERECO', 'Endereço da Empresa'),
            '{{empresa_telefone}}': os.getenv('EMPRESA_TELEFONE', '(00) 0000-0000'),
            '{{empresa_email}}': os.getenv('EMPRESA_EMAIL', 'contato@empresa.com'),
            '{{empresa_site}}': os.getenv('EMPRESA_SITE', 'www.empresa.com'),
            
            # Datas
            '{{data_proposta}}': datetime.now().strftime('%d/%m/%Y'),
            '{{validade_proposta}}': str(o.validade_dias),
        }
    
    def _substituir_texto(self, doc, dados):
        """Substitui as chaves pelos valores mantendo a formatação"""
        
        for paragraph in doc.paragraphs:
            for key, value in dados.items():
                if key in paragraph.text:
                    for run in paragraph.runs:
                        if key in run.text:
                            run.text = run.text.replace(key, str(value))
        
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for key, value in dados.items():
                            if key in paragraph.text:
                                for run in paragraph.runs:
                                    if key in run.text:
                                        run.text = run.text.replace(key, str(value))
    
    def gerar(self, output_path=None):
        """Gera a proposta comercial em DOCX"""
        if not os.path.exists(self.TEMPLATE_PATH):
            raise FileNotFoundError(f"Template não encontrado: {self.TEMPLATE_PATH}")
        
        doc = Document(self.TEMPLATE_PATH)
        dados = self._preparar_dados()
        self._substituir_texto(doc, dados)
        
        if not output_path:
            output_dir = 'media/propostas'
            os.makedirs(output_dir, exist_ok=True)
            output_path = os.path.join(
                output_dir, 
                f'proposta_comercial_{self.orcamento.id}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.docx'
            )
        
        doc.save(output_path)
        return output_path
