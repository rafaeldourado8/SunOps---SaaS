"""
Script para padronizar chaves do PROPOSTA.docx
Converte {{CHAVE_MAIUSCULA}} para {{chave_minuscula}}
"""
from docx import Document
import os

def padronizar_chaves():
    arquivo = 'docs/templates/PROPOSTA_COMERCIAL_TEMPLATE.docx'
    
    print(f"[INFO] Processando {arquivo}...")
    doc = Document(arquivo)
    
    # Mapeamento completo de chaves
    mapeamento = {
        # Cliente
        '{{CLIENTE_NOME}}': '{{cliente_nome}}',
        '{{CLIENTE_ENDERECO}}': '{{cliente_endereco}}',
        '{{CLIENTE_BAIRRO}}': '{{cliente_bairro}}',
        '{{CLIENTE_CIDADE}}': '{{cliente_cidade}}',
        '{{CLIENTE_ESTADO}}': '{{cliente_estado}}',
        '{{CLIENTE_CEP}}': '{{cliente_cep}}',
        '{{CLIENTE_TELEFONE}}': '{{cliente_telefone}}',
        '{{CLIENTE_EMAIL}}': '{{cliente_email}}',
        
        # Sistema
        '{{CONSUMO_MENSAL}}': '{{consumo_mensal}}',
        '{{POTENCIA_TOTAL_KWP}}': '{{potencia_sistema}}',
        '{{PAINEIS_QTD}}': '{{quantidade_paineis}}',
        '{{PAINEIS_POTENCIA}}': '{{potencia_painel}}',
        '{{PAINEIS_MARCA}}': '{{marca_painel}}',
        '{{INVERSOR_POTENCIA}}': '{{potencia_inversor}}',
        '{{INVERSOR_MARCA}}': '{{marca_inversor}}',
        '{{INVERSOR_QTD}}': '{{quantidade_inversores}}',
        '{{GERACAO_ESTIMADA_KWH}}': '{{geracao_mensal}}',
        '{{GERACAO_ANUAL}}': '{{geracao_anual}}',
        
        # Financeiro
        '{{VALOR_FINAL}}': '{{valor_total}}',
        '{{VALOR_TOTAL}}': '{{valor_total}}',
        '{{VALOR_PARCELA}}': '{{valor_parcela}}',
        '{{QUANTIDADE_PARCELAS}}': '{{quantidade_parcelas}}',
        '{{ECONOMIA_MENSAL}}': '{{economia_mensal}}',
        '{{ECONOMIA_ANUAL}}': '{{economia_anual}}',
        '{{PAYBACK}}': '{{payback}}',
        
        # Tecnico
        '{{HSP}}': '{{hsp}}',
        '{{PERDA_SISTEMA}}': '{{perdas_sistema}}',
        '{{DEGRADACAO_ANUAL}}': '{{degradacao_anual}}',
        
        # Empresa
        '{{EMPRESA_NOME}}': '{{empresa_nome}}',
        '{{EMPRESA_CNPJ}}': '{{empresa_cnpj}}',
        '{{EMPRESA_ENDERECO}}': '{{empresa_endereco}}',
        '{{EMPRESA_TELEFONE}}': '{{empresa_telefone}}',
        '{{EMPRESA_EMAIL}}': '{{empresa_email}}',
        '{{EMPRESA_SITE}}': '{{empresa_site}}',
        
        # Datas
        '{{DATA_CRIACAO}}': '{{data_proposta}}',
        '{{DATA_PROPOSTA}}': '{{data_proposta}}',
        '{{VALIDADE_PROPOSTA}}': '{{validade_proposta}}',
    }
    
    contador = 0
    
    # Substituir em paragrafos
    print("[INFO] Processando paragrafos...")
    for paragrafo in doc.paragraphs:
        texto_original = paragrafo.text
        for chave_antiga, chave_nova in mapeamento.items():
            if chave_antiga in paragrafo.text:
                for run in paragrafo.runs:
                    if chave_antiga in run.text:
                        run.text = run.text.replace(chave_antiga, chave_nova)
                        contador += 1
                        print(f"  [OK] {chave_antiga} -> {chave_nova}")
    
    # Substituir em tabelas
    print("[INFO] Processando tabelas...")
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragrafo in cell.paragraphs:
                    for chave_antiga, chave_nova in mapeamento.items():
                        if chave_antiga in paragrafo.text:
                            for run in paragrafo.runs:
                                if chave_antiga in run.text:
                                    run.text = run.text.replace(chave_antiga, chave_nova)
                                    contador += 1
                                    print(f"  [OK] {chave_antiga} -> {chave_nova}")
    
    # Salvar
    print(f"\n[INFO] Salvando arquivo...")
    doc.save(arquivo)
    
    print(f"\n[SUCESSO] {contador} substituicoes realizadas!")
    print(f"[INFO] Arquivo atualizado: {arquivo}")

if __name__ == '__main__':
    padronizar_chaves()
