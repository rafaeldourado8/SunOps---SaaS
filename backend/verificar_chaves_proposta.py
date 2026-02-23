"""
Script para verificar as chaves presentes no PROPOSTA.docx
"""
from docx import Document
import re

def verificar_chaves():
    arquivo = 'docs/templates/PROPOSTA_COMERCIAL_TEMPLATE.docx'
    doc = Document(arquivo)
    
    chaves_encontradas = set()
    
    # Buscar em paragrafos
    for paragrafo in doc.paragraphs:
        matches = re.findall(r'\{\{([^}]+)\}\}', paragrafo.text)
        chaves_encontradas.update(matches)
    
    # Buscar em tabelas
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragrafo in cell.paragraphs:
                    matches = re.findall(r'\{\{([^}]+)\}\}', paragrafo.text)
                    chaves_encontradas.update(matches)
    
    print(f"[INFO] Total de chaves encontradas: {len(chaves_encontradas)}\n")
    print("[CHAVES ENCONTRADAS]")
    for chave in sorted(chaves_encontradas):
        print(f"  - {{{{{chave}}}}}")

if __name__ == '__main__':
    verificar_chaves()
