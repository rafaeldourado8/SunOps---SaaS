"""
Corrigir chaves restantes em MAIUSCULO
"""
from docx import Document

arquivo = 'docs/templates/PROPOSTA_COMERCIAL_TEMPLATE.docx'
doc = Document(arquivo)

mapeamento = {
    '{{INVERSOR_POTENCIA}}': '{{potencia_inversor}}',
    '{{PAINEIS_POTENCIA}}': '{{potencia_painel}}',
    '{{POTENCIA_TOTAL_KWP}}': '{{potencia_sistema}}',
}

contador = 0

for paragrafo in doc.paragraphs:
    for antiga, nova in mapeamento.items():
        if antiga in paragrafo.text:
            for run in paragrafo.runs:
                if antiga in run.text:
                    run.text = run.text.replace(antiga, nova)
                    contador += 1
                    print(f"[OK] {antiga} -> {nova}")

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragrafo in cell.paragraphs:
                for antiga, nova in mapeamento.items():
                    if antiga in paragrafo.text:
                        for run in paragrafo.runs:
                            if antiga in run.text:
                                run.text = run.text.replace(antiga, nova)
                                contador += 1
                                print(f"[OK] {antiga} -> {nova}")

doc.save(arquivo)
print(f"\n[SUCESSO] {contador} substituicoes realizadas!")
