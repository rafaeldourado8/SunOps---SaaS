"""
Script para corrigir o arquivo PROPOSTA.docx com as chaves corretas
"""
from docx import Document
import os

def corrigir_proposta():
    """Corrige as chaves do template PROPOSTA.docx"""
    
    arquivo_entrada = 'PROPOSTA.docx'
    arquivo_saida = 'docs/templates/PROPOSTA_COMERCIAL_TEMPLATE.docx'
    
    if not os.path.exists(arquivo_entrada):
        print(f"[ERRO] Arquivo {arquivo_entrada} nao encontrado!")
        return
    
    print(f"[INFO] Abrindo {arquivo_entrada}...")
    doc = Document(arquivo_entrada)
    
    # Mapeamento de chaves antigas para novas (padronizadas)
    mapeamento = {
        # Cliente
        '{{CLIENTE_NOME}}': '{{cliente_nome}}',
        '{{CLIENTE_ENDERECO}}': '{{cliente_endereco}}',
        '{{CLIENTE_BAIRRO}}': '{{cliente_bairro}}',
        '{{CLIENTE_CIDADE}}': '{{cliente_cidade}}',
        '{{CLIENTE_ESTADO}}': '{{cliente_estado}}',
        '{{CLIENTE_CEP}}': '{{cliente_cep}}',
        '{{CLIENTE_TELEFONE}}': '{{cliente_telefone}}',
        '{{CLIENTE_EMAIL}}': '{{cliente_email}}',
        
        # Sistema
        '{{CONSUMO_MENSAL}}': '{{consumo_mensal}}',
        '{{POTENCIA_TOTAL_KWP}}': '{{potencia_sistema}}',
        '{{PAINEIS_QTD}}': '{{quantidade_paineis}}',
        '{{PAINEIS_POTENCIA}}': '{{potencia_painel}}',
        '{{PAINEIS_MARCA}}': 'CANADIAN SOLAR',
        '{{INVERSOR_POTENCIA}}': '{{potencia_inversor}}',
        '{{INVERSOR_MARCA}}': 'GROWATT',
        '{{GERACAO_ESTIMADA_KWH}}': '{{geracao_mensal}}',
        '{{GERACAO_ANUAL}}': '{{geracao_anual}}',
        
        # Financeiro
        '{{VALOR_TOTAL}}': '{{valor_total}}',
        '{{VALOR_PARCELA}}': '{{valor_parcela}}',
        '{{QUANTIDADE_PARCELAS}}': '{{quantidade_parcelas}}',
        '{{ECONOMIA_MENSAL}}': '{{economia_mensal}}',
        '{{ECONOMIA_ANUAL}}': '{{economia_anual}}',
        '{{PAYBACK}}': '{{payback}}',
        
        # Tecnico
        '{{HSP}}': '{{hsp}}',
        '{{PERDA_SISTEMA}}': '{{perdas_sistema}}',
        
        # Empresa
        '{{EMPRESA_NOME}}': '{{empresa_nome}}',
        '{{EMPRESA_CNPJ}}': '{{empresa_cnpj}}',
        '{{EMPRESA_ENDERECO}}': '{{empresa_endereco}}',
        '{{EMPRESA_TELEFONE}}': '{{empresa_telefone}}',
        '{{EMPRESA_EMAIL}}': '{{empresa_email}}',
        '{{EMPRESA_SITE}}': '{{empresa_site}}',
        
        # Datas
        '{{DATA_CRIACAO}}': '{{data_proposta}}',
        '{{DATA_PROPOSTA}}': '{{data_proposta}}',
        '{{VALIDADE_PROPOSTA}}': '{{validade_proposta}}',
    }
    
    contador = 0
    
    # Substituir em paragrafos
    print("[INFO] Processando paragrafos...")
    for paragrafo in doc.paragraphs:
        for chave_antiga, chave_nova in mapeamento.items():
            if chave_antiga in paragrafo.text:
                for run in paragrafo.runs:
                    if chave_antiga in run.text:
                        run.text = run.text.replace(chave_antiga, chave_nova)
                        contador += 1
                        print(f"  [OK] {chave_antiga} -> {chave_nova}")
    
    # Substituir em tabelas
    print("[INFO] Processando tabelas...")
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragrafo in cell.paragraphs:
                    for chave_antiga, chave_nova in mapeamento.items():
                        if chave_antiga in paragrafo.text:
                            for run in paragrafo.runs:
                                if chave_antiga in run.text:
                                    run.text = run.text.replace(chave_antiga, chave_nova)
                                    contador += 1
                                    print(f"  [OK] {chave_antiga} -> {chave_nova}")
    
    # Criar diretorio se nao existir
    os.makedirs(os.path.dirname(arquivo_saida), exist_ok=True)
    
    # Salvar arquivo corrigido
    print(f"\n[INFO] Salvando {arquivo_saida}...")
    doc.save(arquivo_saida)
    
    print(f"\n[SUCESSO] Concluido! {contador} substituicoes realizadas")
    print(f"[INFO] Arquivo salvo em: {arquivo_saida}")
    print("\n[INFO] Proximos passos:")
    print("1. Fazer upload do arquivo corrigido na plataforma")
    print("2. Testar a geracao de proposta com dados reais")

if __name__ == '__main__':
    corrigir_proposta()
